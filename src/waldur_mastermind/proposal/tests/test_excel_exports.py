import io

from openpyxl import load_workbook
from rest_framework import status

from waldur_mastermind.proposal.tests.test_exports import BaseExportTest, read_csv


def _as_csv(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


class ExcelExportTest(BaseExportTest):
    def workbook_rows(self, response):
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(".xlsx", response["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(response.content)).active
        return sheet, [list(row) for row in sheet.iter_rows(values_only=True)]

    def test_same_rows_as_the_csv(self):
        for action in ("export-proposals", "export-reviews"):
            with self.subTest(action=action):
                csv_rows = read_csv(self.get("staff", action))
                sheet, rows = self.workbook_rows(
                    self.get("staff", action, file_format="xlsx")
                )
                self.assertEqual(rows[0], csv_rows[0])
                self.assertEqual(len(rows), len(csv_rows))
                for row, csv_row in zip(rows[1:], csv_rows[1:]):
                    self.assertEqual([_as_csv(value) for value in row], csv_row)

    def test_numbers_are_numbers_and_formulas_are_text(self):
        self.proposal.name = '=HYPERLINK("https://evil.example","x")'
        self.proposal.save()
        sheet, rows = self.workbook_rows(
            self.get("staff", "export-proposals", file_format="xlsx")
        )
        header = rows[0]
        index = next(i for i, r in enumerate(rows) if r[0] == "ROUND-001")
        row = rows[index]
        self.assertEqual(row[header.index("Name")], self.proposal.name)
        name_cell = sheet.cell(row=index + 1, column=header.index("Name") + 1)
        self.assertEqual(name_cell.data_type, "s")
        self.assertIn(80000, row)

    def test_access_is_the_csv_exports(self):
        response = self.get("user", "export-proposals", file_format="xlsx")
        self.assertIn(
            response.status_code, (status.HTTP_403_FORBIDDEN, status.HTTP_404_NOT_FOUND)
        )

    def test_openpyxl_is_not_loaded_at_startup(self):
        # Loaded by this test module itself; what matters is that the app's
        # modules do not import it at their top.
        from waldur_mastermind.proposal import exports

        self.assertNotIn("openpyxl", vars(exports))
